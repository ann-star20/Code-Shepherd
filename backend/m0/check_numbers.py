import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from pptx import Presentation

PPTX = r"D:\Code-Shepherd\presentation\kms_presentation_v2(수정중).pptx"
prs = Presentation(PPTX)

print(f"총 슬라이드: {len(prs.slides)}\n")
for i, slide in enumerate(prs.slides):
    all_text = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                t = para.text.strip()
                if t: all_text.append(t)

    # 배지(번호 배지), 푸터 번호, 스텝/스테이지 번호 추출
    badge = next((t for t in all_text if re.match(r'^0\d\s+|^1\d\s+', t)), None)
    footer_num = next((t for t in all_text if re.match(r'^\d+\s*/\s*\d+$', t)), None)
    steps = [t for t in all_text if re.match(r'^(Step|Stage|S)\s*\d+', t, re.I)]
    print(f"S{i+1:02d} | badge={badge!r:35s} | footer={footer_num!r:12s} | steps={steps}")
