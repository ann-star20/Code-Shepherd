import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from pptx import Presentation
from pptx.util import Inches, Pt

PPTX = r"D:\Code-Shepherd\presentation\kms_presentation_v2(수정중).pptx"
prs = Presentation(PPTX)

print(f"총 슬라이드: {len(prs.slides)}\n")
print("="*60)

for i, slide in enumerate(prs.slides):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                t = para.text.strip()
                if t:
                    texts.append(t)
    print(f"\n[Slide {i+1}]")
    for t in texts[:25]:   # 최대 25줄만
        print(f"  {repr(t)}")
    if len(texts) > 25:
        print(f"  ... +{len(texts)-25}줄 더")
