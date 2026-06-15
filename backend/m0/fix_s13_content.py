import sys, io, copy
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from pptx import Presentation
from pptx.util import Pt
from lxml import etree

PPTX = r"D:\Code-Shepherd\presentation\kms_presentation_v2(수정중).pptx"
prs  = Presentation(PPTX)
sl   = prs.slides[12]   # S13 (0-based)
shapes = list(sl.shapes)

def set_shape_text(shape, *lines):
    """모든 단락의 텍스트를 교체 — run 서식(폰트·색·사이즈) 보존"""
    tf = shape.text_frame
    # 현재 run 서식 샘플 저장
    sample_run = None
    for para in tf.paragraphs:
        if para.runs:
            sample_run = para.runs[0]
            break

    # 기존 단락 전체 텍스트 교체
    paras = tf.paragraphs
    for i, line in enumerate(lines):
        if i < len(paras):
            p = paras[i]
            # 모든 run 비우고 첫 run에만 텍스트
            if p.runs:
                p.runs[0].text = line
                for r in p.runs[1:]:
                    r.text = ""
            else:
                # run 없는 경우 새로 추가
                from pptx.oxml.ns import qn
                r_elem = copy.deepcopy(sample_run._r) if sample_run else None
                if r_elem is not None:
                    r_elem.find(qn('a:t')).text = line
                    p._p.append(r_elem)
        else:
            # 단락 부족하면 마지막 단락 복사 후 추가
            last_p = paras[-1]._p
            new_p = copy.deepcopy(last_p)
            # 새 단락 텍스트 설정
            from pptx.oxml.ns import qn
            for t_elem in new_p.findall('.//' + qn('a:t')):
                t_elem.text = line
                break
            tf._txBody.append(new_p)
            paras = tf.paragraphs  # 갱신

    # 남은 단락 비우기 (lines 수보다 단락이 많은 경우)
    for i in range(len(lines), len(tf.paragraphs)):
        p = tf.paragraphs[i]
        for r in p.runs:
            r.text = ""

# ────────────────────────────────────────────────────────────
# 교체 작업
# ────────────────────────────────────────────────────────────
changes = [
    # idx  new_lines
    # ── 제목·부제목 ──────────────────────────────────────────
    (1,  ["Graph DB — 코드 리뷰 AI를 위한 외부 지식 레이어"]),
    (2,  ["OWASP · 공식 프레임워크 문서를 Knowledge Graph로 구조화 — AI가 근거 기반으로 판단할 수 있도록 지원"]),

    # ── 좌 패널: 구축 목적 5단계 ─────────────────────────────
    (5,  ["Graph DB 구축 목적"]),
    (7,  ["코드 리뷰 AI 정확도 향상 — 보안·성능 기준에 최신 지식 필요"]),
    (9,  ["사내 KMS + LLM 지식만으론 한계 — 지식 컷오프로 최신 외부 정보 반영 불가"]),
    (11, ["외부 공식 문서 수집 — OWASP · 최신 프레임워크 · Best Practice 실시간 반영"]),
    (13, ["4-Layer 구조화 — Domain → Concept → Mechanism → CodeExample"]),
    (15, ["코드 리뷰 Agent 연계 — 취약점 탐지 시 공식 근거 + 개선 코드 자동 제시"]),

    # ── 우 상단: 4-Layer 구조 예시 ───────────────────────────
    (17, ["4-Layer 구조 활용 예시"]),
    (18, ["  Security  →  SQL Injection  (Domain → Concept)"]),
    (19, ["  └─  Prepared Statement  (Mechanism)"]),
    (20, ["  └─  실제 패치 코드 예제  (CodeExample)"]),

    # ── 우 하단: 사용자 가치 & 확장성 ────────────────────────
    (22, ["사용자 가치 & KMS 확장성"]),
    (23, ["v  기존: '취약합니다'  →  개선: 'OWASP 기준 SQL Injection 취약점,"
          " Prepared Statement 적용 필요'"]),
    (24, ["v  Code Review · Curriculum Agent가 Graph DB 조회",
          "    → 공식 문서 근거 자동 인용"]),
    (25, ["v  KMS 공통 자산화 — 조직 AI 서비스가",
          "    공유하는 외부 지식 레이어로 축적"]),
]

for idx, lines in changes:
    sh = shapes[idx]
    old = " | ".join(p.text.strip() for p in sh.text_frame.paragraphs if p.text.strip())
    set_shape_text(sh, *lines)
    new = " | ".join(p.text.strip() for p in sh.text_frame.paragraphs if p.text.strip())
    print(f"  [shape {idx:2d}] {old!r}")
    print(f"            → {new!r}")

prs.save(PPTX)
print(f"\n저장 완료: {PPTX}")
