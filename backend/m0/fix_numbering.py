import sys, io, re, copy
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from pptx import Presentation
from pptx.util import Pt

SRC  = r"D:\Code-Shepherd\presentation\kms_presentation_v2(수정중).pptx"
OUT  = r"D:\Code-Shepherd\presentation\kms_presentation_v2(수정중).pptx"

prs = Presentation(SRC)

def replace_run_text(shape, old, new):
    """run 서식 유지하며 텍스트만 교체"""
    changed = False
    if not shape.has_text_frame:
        return False
    for para in shape.text_frame.paragraphs:
        full = "".join(r.text for r in para.runs)
        if old in full:
            # 텍스트가 여러 run에 분산되어 있을 수 있으므로
            # 첫 번째 run에 전체 텍스트를 넣고 나머지 비움
            if para.runs:
                para.runs[0].text = full.replace(old, new)
                for r in para.runs[1:]:
                    r.text = ""
            changed = True
    return changed

# ────────────────────────────────────────────────────────────
# 수정 계획 (slide_idx 는 0-based)
#   (slide_idx, old_text, new_text, 설명)
# ────────────────────────────────────────────────────────────
FIXES = [
    # ── S07 (idx=6) : 배지 + 푸터 ──
    (6, "09  데이터 구조 · ERD",      "05  데이터 구조 · ERD",      "S07 배지"),
    (6, "10 / 12",                   "6 / 13",                    "S07 푸터"),

    # ── S08 (idx=7) : 배지 + 푸터 ──
    (7, "05  활용 사례 ①",            "06  활용 사례 ①",            "S08 배지"),
    (7, "6 / 12",                    "7 / 13",                    "S08 푸터"),

    # ── S11 (idx=10) : 배지 + 푸터 ──
    (10, "06  활용 사례 ②",           "07  활용 사례 ②",            "S11 배지"),
    (10, "7 / 12",                   "8 / 13",                    "S11 푸터"),

    # ── S13 (idx=12) : 배지 + 푸터 ──
    (12, "07  Graph DB 실험",         "08  Graph DB 실험",          "S13 배지"),
    (12, "8 / 12",                   "9 / 13",                    "S13 푸터"),

    # ── S15 (idx=14) : 배지 + 푸터 ──
    (14, "08  전체 아키텍처",          "09  전체 아키텍처",           "S15 배지"),
    (14, "9 / 12",                   "10 / 13",                   "S15 푸터"),

    # ── S16 (idx=15) : 배지 + 푸터 ──
    (15, "09  데이터 구조 · ERD",      "10  데이터 구조 · ERD",      "S16 배지"),
    (15, "10 / 12",                  "11 / 13",                   "S16 푸터"),

    # ── S17 (idx=16) : 배지 + 푸터 ──
    (16, "10  시연",                  "11  시연",                   "S17 배지"),
    (16, "11 / 12",                  "12 / 13",                   "S17 푸터"),

    # ── S18 (idx=17) : 배지 + 푸터 ──
    (17, "11  기대 효과 · 마무리",     "12  기대 효과 · 마무리",      "S18 배지"),
    (17, "12 / 12",                  "13 / 13",                   "S18 푸터"),

    # ── S01~S05 : 푸터 총계만 /12 → /13 ──
    (0, "1 / 12",  "1 / 13",  "S01 푸터"),
    (1, "2 / 12",  "2 / 13",  "S02 푸터"),
    (2, "3 / 12",  "3 / 13",  "S03 푸터"),
    (3, "4 / 12",  "4 / 13",  "S04 푸터"),
    (4, "5 / 12",  "5 / 13",  "S05 푸터"),
]

results = []
for (slide_idx, old, new, label) in FIXES:
    slide = prs.slides[slide_idx]
    hit = False
    for shape in slide.shapes:
        if replace_run_text(shape, old, new):
            hit = True
    status = "✓" if hit else "✗ NOT FOUND"
    results.append(f"  [{status}] S{slide_idx+1:02d} {label}: {old!r} → {new!r}")

for r in results:
    print(r)

prs.save(OUT)
print(f"\n저장 완료: {OUT}")
