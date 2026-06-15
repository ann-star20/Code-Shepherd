import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from pptx import Presentation
from pptx.util import Inches

PPTX = r"D:\Code-Shepherd\presentation\kms_presentation_v2(수정중).pptx"
prs  = Presentation(PPTX)
sl   = prs.slides[12]  # S13 (0-based)

print("=== S13 전체 shape 목록 ===")
for i, sh in enumerate(sl.shapes):
    if sh.has_text_frame:
        texts = [p.text.strip() for p in sh.text_frame.paragraphs if p.text.strip()]
        l = round(sh.left   / 914400, 3)
        t = round(sh.top    / 914400, 3)
        w = round(sh.width  / 914400, 3)
        h = round(sh.height / 914400, 3)
        print(f"  [{i:2d}] ({l},{t}) {w}x{h}in  texts={texts}")
