import sys
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation

prs = Presentation("presentation/3차프로젝트_v6.pptx")
slide = prs.slides[11]

print(f"슬라이드 12 도형 수: {len(slide.shapes)}")
for i, shape in enumerate(slide.shapes):
    stype = shape.shape_type
    print(f"\n[{i}] type={stype} name={shape.name!r}")
    if shape.has_text_frame:
        for j, para in enumerate(shape.text_frame.paragraphs):
            text = para.text.strip()
            if text:
                print(f"  para[{j}]: {text!r}")
